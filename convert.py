#!/usr/bin/env python3
"""
Markdown → DOCX converter
Supports: standard Markdown, Obsidian-flavoured Markdown, Mermaid charts
"""

import argparse
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt
from mdit_py_plugins.front_matter import front_matter_plugin
from mdit_py_plugins.tasklists import tasklists_plugin

# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------
QUOTE_BAR = RGBColor(0xD0, 0xD7, 0xDE)     # left border on blockquotes
CALLOUT_COLOURS = {
    "note":      RGBColor(0x08, 0x80, 0xFF),
    "info":      RGBColor(0x08, 0x80, 0xFF),
    "tip":       RGBColor(0x0B, 0xD0, 0x7C),
    "success":   RGBColor(0x0B, 0xD0, 0x7C),
    "important": RGBColor(0xAB, 0x5B, 0xF8),
    "warning":   RGBColor(0xFF, 0xA0, 0x00),
    "caution":   RGBColor(0xFF, 0xA0, 0x00),
    "danger":    RGBColor(0xFF, 0x35, 0x35),
    "failure":   RGBColor(0xFF, 0x35, 0x35),
    "bug":       RGBColor(0xFF, 0x35, 0x35),
    "example":   RGBColor(0x78, 0x4C, 0xFF),
    "quote":     RGBColor(0x9E, 0xA3, 0xAD),
}
HIGHLIGHT_BG = RGBColor(0xFF, 0xFF, 0x00)  # ==highlight==
TABLE_HEADER_BG = RGBColor(0xF0, 0xF2, 0xF4)
TABLE_BORDER = RGBColor(0xCC, 0xD0, 0xD5)


# ---------------------------------------------------------------------------
# XML helpers
# ---------------------------------------------------------------------------

def _set_cell_bg(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
    tcPr.append(shd)


def _set_cell_border(cell, **edges):
    """edges: top/bottom/left/right with (color_hex, sz, space, val) tuples."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge, (color, sz, space, val) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), str(space))
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _para_shading(para, rgb: RGBColor):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
    pPr.append(shd)


def _para_left_border(para, rgb: RGBColor, sz: int = 24):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
    pBdr.append(left)
    pPr.append(pBdr)


def _run_shading(run, rgb: RGBColor):
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}")
    rPr.append(shd)


def _set_para_spacing(para, before: int = 0, after: int = 0, line: Optional[int] = None):
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    if line is not None:
        spacing.set(qn("w:line"), str(line))
        spacing.set(qn("w:lineRule"), "auto")
    pPr.append(spacing)


def _hyperlink(para, text: str, url: str):
    """Add a hyperlink run to a paragraph."""
    part = para.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "Hyperlink")
    rPr.append(rStyle)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    new_run.append(t)
    hyperlink.append(new_run)
    para._p.append(hyperlink)


# ---------------------------------------------------------------------------
# Mermaid rendering
# ---------------------------------------------------------------------------

def render_mermaid(code: str, tmp_dir: str) -> Optional[str]:
    """Render a mermaid diagram to a PNG file. Returns the path or None."""
    mmd_file = os.path.join(tmp_dir, "diagram.mmd")
    png_file = os.path.join(tmp_dir, "diagram.png")
    with open(mmd_file, "w") as f:
        f.write(code)
    cfg = os.path.join(tmp_dir, "puppeteer.json")
    with open(cfg, "w") as f:
        f.write('{"args": ["--no-sandbox", "--disable-setuid-sandbox"]}')
    try:
        result = subprocess.run(
            ["mmdc", "-i", mmd_file, "-o", png_file, "-p", cfg, "-b", "white"],
            capture_output=True, text=True, timeout=30
        )
        if result.returncode == 0 and os.path.exists(png_file):
            return png_file
        print(f"  [mermaid warning] {result.stderr.strip()}", file=sys.stderr)
        return None
    except (subprocess.TimeoutExpired, FileNotFoundError) as e:
        print(f"  [mermaid warning] {e}", file=sys.stderr)
        return None


# ---------------------------------------------------------------------------
# Obsidian pre-processing
# ---------------------------------------------------------------------------

def preprocess_obsidian(text: str) -> tuple[str, dict]:
    """
    Transform Obsidian-specific syntax into something markdown-it understands,
    or into special markers we can detect during rendering.
    Strips all hidden / non-printing content before conversion.
    Returns (transformed_text, metadata_dict).
    """
    # --- Strip YAML frontmatter (--- delimiters) ---
    meta = {}
    fm_match = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.DOTALL)
    if fm_match:
        for line in fm_match.group(1).splitlines():
            if ":" in line:
                k, _, v = line.partition(":")
                meta[k.strip()] = v.strip()
        text = text[fm_match.end():]

    # --- Strip TOML frontmatter (+++ delimiters) ---
    toml_match = re.match(r"^\+\+\+\s*\n(.*?)\n\+\+\+\s*\n", text, re.DOTALL)
    if toml_match:
        text = text[toml_match.end():]

    # --- Strip HTML comments <!-- ... --> ---
    text = re.sub(r"<!--.*?-->", "", text, flags=re.DOTALL)

    # --- Strip <script> blocks ---
    text = re.sub(r"<script[^>]*>.*?</script>", "", text, flags=re.DOTALL | re.IGNORECASE)

    # --- Strip <style> blocks ---
    text = re.sub(r"<style[^>]*>.*?</style>", "", text, flags=re.DOTALL | re.IGNORECASE)

    # --- Strip <details>/<summary> wrapper tags, keep inner content ---
    text = re.sub(r"<summary[^>]*>.*?</summary>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"</?details[^>]*>", "", text, flags=re.IGNORECASE)

    # --- Strip zero-width and other invisible Unicode characters ---
    # Zero-width space (U+200B), zero-width non-joiner (U+200C),
    # zero-width joiner (U+200D), word joiner (U+2060), BOM (U+FEFF),
    # soft hyphen (U+00AD), left/right marks (U+200E/F)
    text = re.sub(r"[­​‌‍‎‏⁠﻿]", "", text)

    # --- Strip Obsidian comments %%...%% ---
    text = re.sub(r"%%.*?%%", "", text, flags=re.DOTALL)

    # --- ==highlights== → <mark>text</mark> (HTML passthrough) ---
    text = re.sub(r"==(.+?)==", r"<mark>\1</mark>", text)

    # --- Obsidian callouts: > [!TYPE] → blockquote with CALLOUT: prefix ---
    def replace_callout(m):
        kind = m.group(1).lower()
        title = m.group(2).strip() if m.group(2) else kind.upper()
        return f"> CALLOUT:{kind}:{title}\n"
    text = re.sub(r"^> \[!(\w+)\][ \t]*(.*)?$", replace_callout, text, flags=re.MULTILINE)

    # --- Wikilinks [[Page|alias]] and [[Page]] ---
    # Use angle-bracket URLs so spaces are valid CommonMark
    text = re.sub(r"\[\[([^\]|]+)\|([^\]]+)\]\]", lambda m: f"[{m.group(2)}](<{m.group(1)}>)", text)
    text = re.sub(r"\[\[([^\]]+)\]\]", lambda m: f"[{m.group(1)}](<{m.group(1)}>)", text)

    # --- Embedded images ![[file.png]] ---
    text = re.sub(r"!\[\[([^\]]+)\]\]", r"![\1](\1)", text)

    # --- Inline math $...$ → `...` (keep readable without LaTeX) ---
    text = re.sub(r"(?<!\$)\$(?!\$)(.+?)(?<!\$)\$(?!\$)", r"`\1`", text)
    # Block math $$...$$ → fenced code block
    text = re.sub(r"\$\$(.+?)\$\$", r"\n```math\n\1\n```\n", text, flags=re.DOTALL)

    # --- Obsidian tags #tag (not headings) – render as inline code ---
    # Match #tag preceded by whitespace; headings use "# text" (space after #)
    text = re.sub(r"(\s)#([A-Za-z][\w/-]*)", r"\1`#\2`", text)

    return text, meta


# ---------------------------------------------------------------------------
# Inline token renderer
# ---------------------------------------------------------------------------

class InlineRenderer:
    """Renders a list of markdown-it inline tokens onto a python-docx paragraph."""

    def __init__(self, para, doc: Document):
        self.para = para
        self.doc = doc
        self._bold = False
        self._italic = False
        self._strike = False
        self._code = False
        self._mark = False  # ==highlight==

    def render(self, tokens):
        i = 0
        while i < len(tokens):
            tok = tokens[i]
            if tok.type == "text":
                self._text(tok.content)
            elif tok.type == "softbreak":
                self._add_run(" ")
            elif tok.type == "hardbreak":
                self._add_run("\n")
            elif tok.type == "strong_open":
                self._bold = True
            elif tok.type == "strong_close":
                self._bold = False
            elif tok.type == "em_open":
                self._italic = True
            elif tok.type == "em_close":
                self._italic = False
            elif tok.type == "s_open":
                self._strike = True
            elif tok.type == "s_close":
                self._strike = False
            elif tok.type == "code_inline":
                self._code_run(tok.content)
            elif tok.type == "html_inline":
                self._html_inline(tok.content)
            elif tok.type == "link_open":
                href = dict(tok.attrs or {}).get("href", "")
                # Gather link text from next text token
                text_parts = []
                i += 1
                while i < len(tokens) and tokens[i].type != "link_close":
                    if tokens[i].type == "text":
                        text_parts.append(tokens[i].content)
                    i += 1
                link_text = "".join(text_parts) or href
                try:
                    _hyperlink(self.para, link_text, href)
                except Exception:
                    self._add_run(link_text)
            elif tok.type == "image":
                alt = dict(tok.attrs or {}).get("alt", "")
                src = dict(tok.attrs or {}).get("src", "")
                self._inline_image(src, alt)
            i += 1

    def _html_inline(self, html: str):
        mark_open = re.match(r"<mark>", html)
        mark_close = re.match(r"</mark>", html)
        if mark_open:
            self._mark = True
        elif mark_close:
            self._mark = False
        # ignore other HTML

    def _text(self, content: str):
        run = self.para.add_run(content)
        self._style_run(run)

    def _add_run(self, text: str):
        run = self.para.add_run(text)
        self._style_run(run)

    def _style_run(self, run):
        run.bold = self._bold
        run.italic = self._italic
        run.font.strike = self._strike
        # Explicitly unset Word's hidden-text property so nothing is ever invisible
        run.font.hidden = False
        if self._mark:
            run.font.highlight_color = None  # python-docx highlight enum doesn't do custom RGB
            _run_shading(run, HIGHLIGHT_BG)

    def _code_run(self, content: str):
        run = self.para.add_run(content)
        run.font.name = "Consolas"
        run.font.size = Pt(10)

    def _inline_image(self, src: str, alt: str):
        if src and os.path.isfile(src):
            try:
                run = self.para.add_run()
                run.add_picture(src, width=Inches(3))
                return
            except Exception:
                pass
        run = self.para.add_run(f"[image: {alt or src}]")
        run.italic = True


# ---------------------------------------------------------------------------
# Main converter
# ---------------------------------------------------------------------------

class MarkdownToDocx:
    def __init__(self):
        self.doc = Document()
        self._setup_document()
        self.md = self._build_parser()
        self._tmp_dir: Optional[str] = None
        self._list_level = 0
        self._ordered = False

    # ------------------------------------------------------------------
    # Document & style setup
    # ------------------------------------------------------------------

    def _setup_document(self):
        doc = self.doc
        # Page margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1.15)
            section.right_margin = Inches(1.15)

        # Normal paragraph style
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)

        # Heading styles
        heading_sizes = {1: 20, 2: 16, 3: 13, 4: 12, 5: 11, 6: 11}
        heading_colours = {
            1: RGBColor(0x1A, 0x1A, 0x2E),
            2: RGBColor(0x16, 0x21, 0x3E),
            3: RGBColor(0x0F, 0x3A, 0x60),
            4: RGBColor(0x1A, 0x4A, 0x72),
            5: RGBColor(0x1A, 0x4A, 0x72),
            6: RGBColor(0x1A, 0x4A, 0x72),
        }
        for lvl in range(1, 7):
            style = doc.styles[f"Heading {lvl}"]
            style.font.name = "Calibri"
            style.font.size = Pt(heading_sizes[lvl])
            style.font.color.rgb = heading_colours[lvl]
            style.font.bold = True
            style.font.italic = False
            style.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.keep_with_next = True

        # Block Quote style (we create our own)
        try:
            bq = doc.styles.add_style("BlockQuote", 1)
        except Exception:
            bq = doc.styles["BlockQuote"]
        bq.base_style = doc.styles["Normal"]
        bq.font.name = "Calibri"
        bq.font.size = Pt(11)
        bq.font.italic = True
        bq.font.color.rgb = RGBColor(0x55, 0x60, 0x6E)
        bq.paragraph_format.left_indent = Inches(0.3)
        bq.paragraph_format.space_after = Pt(4)

        # Code Block style
        try:
            cb = doc.styles.add_style("CodeBlock", 1)
        except Exception:
            cb = doc.styles["CodeBlock"]
        cb.base_style = doc.styles["Normal"]
        cb.font.name = "Consolas"
        cb.font.size = Pt(9.5)
        cb.paragraph_format.left_indent = Inches(0.2)
        cb.paragraph_format.right_indent = Inches(0.2)
        cb.paragraph_format.space_before = Pt(0)
        cb.paragraph_format.space_after = Pt(0)

    def _build_parser(self) -> MarkdownIt:
        parser = (
            MarkdownIt("commonmark", {"html": True, "typographer": True})
            .enable("strikethrough")
            .enable("table")
        )
        front_matter_plugin(parser)
        tasklists_plugin(parser)
        return parser

    # ------------------------------------------------------------------
    # Public entry point
    # ------------------------------------------------------------------

    def convert(self, md_text: str, output_path: str):
        with tempfile.TemporaryDirectory() as tmp:
            self._tmp_dir = tmp
            text, meta = preprocess_obsidian(md_text)
            if meta.get("title"):
                self.doc.core_properties.title = meta["title"]
            tokens = self.md.parse(text)
            self._walk(tokens)
            self.doc.save(output_path)
            print(f"Saved: {output_path}")

    # ------------------------------------------------------------------
    # Token walker
    # ------------------------------------------------------------------

    def _walk(self, tokens):
        i = 0
        while i < len(tokens):
            tok = tokens[i]

            if tok.type == "front_matter":
                i += 1; continue

            elif tok.type in ("heading_open",):
                level = int(tok.tag[1])
                i += 1  # move to inline
                inline = tokens[i]
                para = self.doc.add_heading("", level=level)
                para.clear()
                InlineRenderer(para, self.doc).render(inline.children or [])
                i += 1  # skip heading_close

            elif tok.type == "paragraph_open":
                i += 1
                inline = tokens[i]
                para = self.doc.add_paragraph(style="Normal")
                _set_para_spacing(para, before=0, after=120)
                InlineRenderer(para, self.doc).render(inline.children or [])
                i += 1  # skip paragraph_close

            elif tok.type == "fence":
                self._add_fence(tok)

            elif tok.type == "code_block":
                self._add_code_lines(tok.content)

            elif tok.type == "hr":
                self._add_hr()

            elif tok.type == "bullet_list_open":
                i, _ = self._walk_list(tokens, i, ordered=False)

            elif tok.type == "ordered_list_open":
                i, _ = self._walk_list(tokens, i, ordered=True)

            elif tok.type == "blockquote_open":
                i = self._walk_blockquote(tokens, i)

            elif tok.type == "table_open":
                i = self._walk_table(tokens, i)

            elif tok.type == "html_block":
                # Try to render <img> tags, otherwise skip
                self._handle_html_block(tok.content)

            i += 1

    # ------------------------------------------------------------------
    # List handling
    # ------------------------------------------------------------------

    def _walk_list(self, tokens, start: int, ordered: bool, level: int = 0) -> tuple[int, None]:
        i = start + 1
        counter = 1
        open_tag = "bullet_list_open" if not ordered else "ordered_list_open"
        close_tag = "bullet_list_close" if not ordered else "ordered_list_close"

        while i < len(tokens):
            tok = tokens[i]
            if tok.type == close_tag:
                return i, None
            elif tok.type == "list_item_open":
                i, counter = self._walk_list_item(tokens, i, ordered, level, counter)
            elif tok.type in (open_tag,):
                i, _ = self._walk_list(tokens, i, ordered, level + 1)
            elif tok.type in ("bullet_list_open", "ordered_list_open"):
                i, _ = self._walk_list(tokens, i, tok.type == "ordered_list_open", level + 1)
            i += 1
        return i, None

    def _walk_list_item(self, tokens, start: int, ordered: bool, level: int, counter: int) -> tuple[int, int]:
        i = start + 1
        while i < len(tokens):
            tok = tokens[i]
            if tok.type == "list_item_close":
                return i, counter + 1
            elif tok.type == "paragraph_open":
                i += 1
                inline = tokens[i]
                bullet = f"{counter}." if ordered else ("•", "◦", "▪")[min(level, 2)]
                indent = Inches(0.25 * (level + 1))
                para = self.doc.add_paragraph(style="Normal")
                para.paragraph_format.left_indent = indent
                para.paragraph_format.first_line_indent = Inches(-0.2)
                _set_para_spacing(para, before=0, after=60)
                label_run = para.add_run(f"{bullet}  ")
                label_run.bold = ordered
                InlineRenderer(para, self.doc).render(inline.children or [])
                i += 1  # skip paragraph_close
            elif tok.type == "bullet_list_open":
                i, _ = self._walk_list(tokens, i, False, level + 1)
            elif tok.type == "ordered_list_open":
                i, _ = self._walk_list(tokens, i, True, level + 1)
            i += 1
        return i, counter + 1

    # ------------------------------------------------------------------
    # Blockquote
    # ------------------------------------------------------------------

    def _walk_blockquote(self, tokens, start: int) -> int:
        i = start + 1
        callout_color: Optional[RGBColor] = None
        first_para = True
        while i < len(tokens):
            tok = tokens[i]
            if tok.type == "blockquote_close":
                return i
            elif tok.type == "paragraph_open":
                i += 1
                inline = tokens[i]
                raw = self._inline_plain(inline.children or [])

                if first_para and raw.startswith("CALLOUT:"):
                    first_para = False
                    parts = raw.split(":", 2)
                    kind = parts[1] if len(parts) > 1 else "note"
                    title = parts[2].strip() if len(parts) > 2 else kind.upper()
                    callout_color = CALLOUT_COLOURS.get(kind.lower(), CALLOUT_COLOURS["note"])
                    icons = {
                        "note": "ℹ", "info": "ℹ", "tip": "💡", "success": "✅",
                        "important": "❗", "warning": "⚠", "caution": "⚠",
                        "danger": "🚫", "failure": "✗", "bug": "🐛",
                        "example": "📌", "quote": "❝",
                    }
                    icon = icons.get(kind.lower(), "›")
                    title_para = self.doc.add_paragraph(style="Normal")
                    _para_left_border(title_para, callout_color, sz=32)
                    title_para.paragraph_format.left_indent = Inches(0.35)
                    _set_para_spacing(title_para, before=80, after=0)
                    run = title_para.add_run(f"{icon}  {kind.upper()}: {title}")
                    run.bold = True
                    run.font.color.rgb = callout_color
                    run.font.size = Pt(10.5)
                else:
                    first_para = False
                    color = callout_color or QUOTE_BAR
                    para = self.doc.add_paragraph(style="BlockQuote")
                    _para_left_border(para, color, sz=32 if callout_color else 24)
                    para.paragraph_format.left_indent = Inches(0.45 if callout_color else 0.35)
                    _set_para_spacing(para, before=0, after=4)
                    InlineRenderer(para, self.doc).render(inline.children or [])
                i += 1
            elif tok.type == "blockquote_open":
                i = self._walk_blockquote(tokens, i)
            i += 1
        return i

    # ------------------------------------------------------------------
    # Fenced code / mermaid
    # ------------------------------------------------------------------

    def _add_fence(self, tok):
        lang = (tok.info or "").strip().lower().split()[0] if tok.info else ""
        content = tok.content.rstrip("\n")

        if lang == "mermaid":
            png = render_mermaid(content, self._tmp_dir)
            if png:
                para = self.doc.add_paragraph()
                _set_para_spacing(para, before=80, after=80)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run()
                run.add_picture(png, width=Inches(5.5))
                caption = self.doc.add_paragraph(style="Normal")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run("Mermaid diagram")
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
                return
            else:
                # Fallback: render as code block with notice
                notice = self.doc.add_paragraph(style="Normal")
                run = notice.add_run("⚠ Mermaid diagram (could not render):")
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xAA, 0x66, 0x00)

        self._add_code_lines(content, lang=lang)

    def _add_code_lines(self, content: str, lang: str = ""):
        lines = content.splitlines() or [""]
        for idx, line in enumerate(lines):
            para = self.doc.add_paragraph(style="CodeBlock")
            if idx == 0:
                _set_para_spacing(para, before=80, after=0)
            elif idx == len(lines) - 1:
                _set_para_spacing(para, before=0, after=80)
            else:
                _set_para_spacing(para, before=0, after=0)
            run = para.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)

    # ------------------------------------------------------------------
    # Table
    # ------------------------------------------------------------------

    def _walk_table(self, tokens, start: int) -> int:
        # Collect all rows
        rows: list[list[list]] = []      # rows[row][col] = list of inline tokens
        in_head = False
        header_rows: set[int] = set()
        cur_row: list[list] = []
        cur_cell: list = []
        alignments: list[str] = []

        i = start + 1
        while i < len(tokens):
            tok = tokens[i]
            if tok.type == "table_close":
                break
            elif tok.type == "thead_open":
                in_head = True
            elif tok.type == "thead_close":
                in_head = False
            elif tok.type == "tr_open":
                cur_row = []
                if in_head:
                    header_rows.add(len(rows))
            elif tok.type == "tr_close":
                rows.append(cur_row)
            elif tok.type in ("th_open", "td_open"):
                align = dict(tok.attrs or {}).get("style", "")
                if "center" in align:
                    alignments.append("center")
                elif "right" in align:
                    alignments.append("right")
                else:
                    alignments.append("left")
                cur_cell = []
            elif tok.type in ("th_close", "td_close"):
                cur_row.append(cur_cell)
            elif tok.type == "inline":
                cur_cell = tok.children or []
            i += 1

        if not rows:
            return i

        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.style = "Table Grid"

        for r_idx, row in enumerate(rows):
            tr = table.rows[r_idx]
            for c_idx, cell_tokens in enumerate(row):
                if c_idx >= ncols:
                    break
                cell = tr.cells[c_idx]
                cell.paragraphs[0].clear()
                para = cell.paragraphs[0]
                _set_para_spacing(para, before=40, after=40)

                if r_idx in header_rows:
                    _set_cell_bg(cell, TABLE_HEADER_BG)

                align = alignments[c_idx] if c_idx < len(alignments) else "left"
                if align == "center":
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif align == "right":
                    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

                ir = InlineRenderer(para, self.doc)
                if r_idx in header_rows:
                    ir._bold = True
                ir.render(cell_tokens)

        # Style table borders
        for row in table.rows:
            for cell in row.cells:
                hex_c = f"{TABLE_BORDER[0]:02X}{TABLE_BORDER[1]:02X}{TABLE_BORDER[2]:02X}"
                _set_cell_border(
                    cell,
                    top=(hex_c, 4, 0, "single"),
                    bottom=(hex_c, 4, 0, "single"),
                    left=(hex_c, 4, 0, "single"),
                    right=(hex_c, 4, 0, "single"),
                )

        # Add some space after the table
        self.doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(2)
        return i

    # ------------------------------------------------------------------
    # HR
    # ------------------------------------------------------------------

    def _add_hr(self):
        para = self.doc.add_paragraph()
        _set_para_spacing(para, before=40, after=40)
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "CCCCCC")
        pBdr.append(bottom)
        pPr.append(pBdr)

    # ------------------------------------------------------------------
    # HTML block
    # ------------------------------------------------------------------

    def _handle_html_block(self, html: str):
        img = re.search(r'<img[^>]+src=["\']([^"\']+)["\']', html)
        if img:
            src = img.group(1)
            if os.path.isfile(src):
                para = self.doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run()
                try:
                    run.add_picture(src, width=Inches(5))
                except Exception:
                    para.add_run(f"[image: {src}]")

    # ------------------------------------------------------------------
    # Utility: extract plain text from inline tokens
    # ------------------------------------------------------------------

    @staticmethod
    def _inline_plain(tokens) -> str:
        return "".join(t.content for t in tokens if t.type == "text")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Convert Markdown (standard + Obsidian) to DOCX",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python convert.py notes.md
  python convert.py notes.md -o report.docx
  python convert.py *.md -o combined.docx
""",
    )
    parser.add_argument("inputs", nargs="+", metavar="FILE", help="Markdown file(s) to convert")
    parser.add_argument("-o", "--output", metavar="OUTPUT", help="Output DOCX path (default: <input>.docx)")
    args = parser.parse_args()

    # Expand globs / validate inputs
    input_paths = []
    for pattern in args.inputs:
        p = Path(pattern)
        if p.exists():
            input_paths.append(p)
        else:
            from glob import glob
            matched = glob(pattern)
            if matched:
                input_paths.extend(Path(m) for m in sorted(matched))
            else:
                print(f"Warning: no file matched '{pattern}'", file=sys.stderr)

    if not input_paths:
        print("Error: no input files found.", file=sys.stderr)
        sys.exit(1)

    if len(input_paths) == 1:
        # Single file
        src = input_paths[0]
        out = Path(args.output) if args.output else src.with_suffix(".docx")
        converter = MarkdownToDocx()
        converter.convert(src.read_text(encoding="utf-8"), str(out))
    else:
        # Multiple files → one combined DOCX (or separate if no -o given)
        if args.output:
            combined_text = ""
            for p in input_paths:
                combined_text += p.read_text(encoding="utf-8") + "\n\n---\n\n"
            converter = MarkdownToDocx()
            converter.convert(combined_text, args.output)
        else:
            for p in input_paths:
                out = p.with_suffix(".docx")
                converter = MarkdownToDocx()
                converter.convert(p.read_text(encoding="utf-8"), str(out))


if __name__ == "__main__":
    main()
